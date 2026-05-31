from openpyxl import load_workbook

class Excelfunction:
    def __init__(self, file_name, sheet_name):
        self.file_name = file_name
        self.sheet_name = sheet_name
        self.workbook = load_workbook(self.file_name)
        self.sheet = self.workbook[self.sheet_name]

    # Read data from Excel
    def get_data_from_excel(self):
        rows = self.sheet.max_row
        columns = self.sheet.max_column
        data = []

        for row in range(2, rows + 1): #Skipping header row
            row_data = []
            for column in range(1, columns + 1):
                value = self.sheet.cell(row=row, column=column).value
                if value is None:
                    value="" # Explicitly handling nulls as empty strings
                row_data.append(value)
            data.append(row_data)

        return data

